# -*- coding: utf-8 -*-
"""
剪映标记读取与导出（Excel）。
  从剪映草稿的复合片段 subdraft/*/draft_content.json（明文）中读取 time_marks，
  按颜色 + 标题配对，生成 Excel（句首时延、句尾时延、汇总，均为公式）。

标记约定（可按需修改 COLOR_KEY/CYAN 常量适配自己的配色）
  句首：青色 "S{n}"  = 音频句首 a      红色  "S{n}L" = 环境A(左) b     橙色 "S{n}R" = 环境B(右) c
  句尾：青色 "T{n}"  = 音频句尾 ae     黄绿色 "T{n}…" = 环境A句尾 be    紫色 "T{n}…" = 环境B句尾 ce
  青色标记按标题首字母 S/T 区分句首/句尾；其余颜色按颜色区分，标题里的数字用于配对，
  无数字的标记按时间归入前一个同类青色标记所在的句子。

前置操作（剪映内，仅一次）：选中带标记的视频片段 → 右键「新建复合片段」→ Ctrl+S 保存。
之后每次改标记后 Ctrl+S，再运行本工具即可读到最新标记。
"""
import glob
import json
import os
import re
import time
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FPS = 60  # 时间码帧率（与剪映提示框显示一致）

ROOT_META = Path(os.environ.get("LOCALAPPDATA", "")) / "JianyingPro/User Data/Projects/com.lveditor.draft/root_meta_info.json"
DEFAULT_ROOTS = [
    "D:/Program Files/Jianying/JianyingPro Drafts",
    os.path.expanduser("~/JianyingPro Drafts"),
]

CYAN = "#00c1cd"
COLOR_KEY = {
    "#fc7265": "b", "#f09b43": "c",
    "#a2c15e": "be", "#a475d2": "ce", "#8d66b3": "ce",
}
KEYS = {
    "a":  ("青色S", "",  "start"), "b":  ("红色",   "L", "start"), "c":  ("橙色", "R", "start"),
    "ae": ("青色T", "",  "end"),   "be": ("黄绿色", "L", "end"),   "ce": ("紫色", "R", "end"),
}
GROUPS = {
    "start": {"sheet": "句首时延", "prefix": "S", "cols": ["a", "b", "c"],
              "hdr": ["a 青色(音频句首)", "b 红色(环境A/左)", "c 橙色(环境B/右)"], "diff": ["b-a", "c-a"]},
    "end":   {"sheet": "句尾时延", "prefix": "T", "cols": ["ae", "be", "ce"],
              "hdr": ["a' 青色(音频句尾)", "b' 黄绿(环境A/左)", "c' 紫色(环境B/右)"], "diff": ["b'-a'", "c'-a'"]},
}
FILL_HEX = {"a": "CCF2F5", "b": "FDDAD5", "c": "FBE3C8", "ae": "CCF2F5", "be": "E6EFCF", "ce": "E4D8F2", "d": "E8E8E8"}

RE_NUM = re.compile(r"[sStT]\s*0*(\d+)")
RE_T = re.compile(r"^\s*[tT]")
RE_SUFFIX = re.compile(r"([lLrR])\s*$")


# ---------------------------------------------------------------- 查找文件

def draft_roots(extra_root=None):
    roots = set()
    if extra_root:
        roots.add(extra_root)
    if ROOT_META.exists():
        try:
            meta = json.load(open(ROOT_META, encoding="utf-8"))
            for p in meta.get("all_draft_store", []):
                fold = p.get("draft_fold_path")
                if fold:
                    roots.add(str(Path(fold).parent))
        except Exception:
            pass
    roots.update(DEFAULT_ROOTS)
    return [r for r in roots if r and os.path.isdir(r)]


def collect_marks(obj, out):
    if isinstance(obj, dict):
        if "mark_items" in obj and isinstance(obj["mark_items"], list):
            out.extend(obj["mark_items"])
        for v in obj.values():
            collect_marks(v, out)
    elif isinstance(obj, list):
        for v in obj:
            collect_marks(v, out)


def find_candidates(draft_filter=None, extra_root=None):
    cands = []
    for root in draft_roots(extra_root):
        for f in glob.glob(os.path.join(root, "*", "subdraft", "*", "draft_content.json")):
            draft_name = Path(f).parents[2].name
            if draft_filter and draft_filter not in draft_name:
                continue
            try:
                with open(f, "rb") as fh:
                    head = fh.read(1)
                if head not in (b"{", b"["):
                    continue  # 加密文件跳过
                data = json.load(open(f, encoding="utf-8"))
            except Exception:
                continue
            items = []
            collect_marks(data, items)
            if items:
                cands.append({"file": f, "draft": draft_name, "mtime": os.path.getmtime(f),
                              "items": items, "n": len(items)})
    cands.sort(key=lambda c: c["mtime"], reverse=True)
    return cands


# ---------------------------------------------------------------- 解析

def sec_to_tc(sec):
    total_frames = round(sec * FPS)
    f = total_frames % FPS
    s = total_frames // FPS
    return f"{s // 3600:02d}:{(s % 3600) // 60:02d}:{s % 60:02d}:{f:02d}"


def classify(color, title):
    color = color.lower()
    if color == CYAN:
        return "ae" if RE_T.match(title) else "a"
    return COLOR_KEY.get(color)


def parse_items(items):
    """返回 rows(dict n -> {key: (sec, title)}), anomalies(list[str]), ignored(list)"""
    marks, ignored, anomalies = [], [], []
    for it in items:
        color = it.get("color", "").lower()
        sec = it["time_range"]["start"] / 1e6
        title = (it.get("title") or "").strip()
        key = classify(color, title)
        if key:
            marks.append({"key": key, "name": KEYS[key][0], "sec": sec, "title": title})
        else:
            ignored.append((color, sec, title))
    marks.sort(key=lambda m: m["sec"])

    rows = defaultdict(dict)
    unnumbered = []
    for m in marks:
        mnum = RE_NUM.search(m["title"])
        if not mnum:
            unnumbered.append(m)
            continue
        n = int(mnum.group(1))
        key = m["key"]
        expect = KEYS[key][1]
        suf = RE_SUFFIX.search(m["title"])
        suf = suf.group(1).upper() if suf else ""
        if expect and suf and suf != expect:
            anomalies.append(f"{GROUPS[KEYS[key][2]]['prefix']}{n} {m['name']} 标题“{m['title']}”后缀与颜色不符（颜色为准）")
        if key in rows[n]:
            anomalies.append(f"{GROUPS[KEYS[key][2]]['prefix']}{n} {m['name']} 重复：{rows[n][key][1]} @ {sec_to_tc(rows[n][key][0])} "
                             f"与 “{m['title']}” @ {sec_to_tc(m['sec'])}，保留前者")
            continue
        rows[n][key] = (m["sec"], m["title"])

    for m in unnumbered:
        key = m["key"]
        anchor = GROUPS[KEYS[key][2]]["cols"][0]
        prefix = GROUPS[KEYS[key][2]]["prefix"]
        anchors = sorted((v[anchor][0], n) for n, v in rows.items() if anchor in v)
        n = None
        for sec, k in anchors:
            if sec <= m["sec"]:
                n = k
        if n is None or key in rows[n]:
            anomalies.append(f"无法归属：{m['name']} “{m['title'] or '(无名)'}” @ {sec_to_tc(m['sec'])}")
            continue
        rows[n][key] = (m["sec"], m["title"] or "(无名)")
        anomalies.append(f"{prefix}{n} {m['name']} 无编号标题“{m['title'] or '(无名)'}”，按时间归入 {prefix}{n}")

    for n in sorted(rows):
        r = rows[n]
        for g in GROUPS.values():
            if not any(k in r for k in g["cols"]):
                continue
            base = g["cols"][0]
            for k in g["cols"]:
                if k not in r:
                    anomalies.append(f"{g['prefix']}{n} 缺少{KEYS[k][0]}标记")
            if base in r:
                for k in g["cols"][1:]:
                    if k in r and r[k][0] < r[base][0]:
                        anomalies.append(f"{g['prefix']}{n} {KEYS[k][0]}早于{KEYS[base][0]}（差值为负）")
    return rows, anomalies, ignored


# ---------------------------------------------------------------- Excel

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center")
BOLD = Font(bold=True)


def fill(key):
    return PatternFill("solid", fgColor=FILL_HEX[key])


def build_group_sheet(ws, g, rows):
    cols = g["cols"]
    hdr1 = [f"{g['prefix']}号"]
    for h in g["hdr"]:
        hdr1 += [h, None]
    hdr1 += g["diff"] + [f"标记名称 ({' / '.join(cols)})"]
    hdr2 = [""] + ["时间码", "秒"] * 3 + ["秒", "秒", ""]
    ws.append(hdr1); ws.append(hdr2)
    for c1, c2 in ((2, 3), (4, 5), (6, 7)):
        ws.merge_cells(start_row=1, start_column=c1, end_row=1, end_column=c2)
    ws.merge_cells("A1:A2"); ws.merge_cells("J1:J2")
    col_fill = {2: cols[0], 3: cols[0], 4: cols[1], 5: cols[1], 6: cols[2], 7: cols[2], 8: "d", 9: "d"}
    for r in (1, 2):
        for c in range(1, 11):
            cell = ws.cell(r, c); cell.font = BOLD; cell.alignment = CENTER; cell.border = BORDER
            if c in col_fill:
                cell.fill = fill(col_fill[c])

    FIRST = 3
    r = FIRST
    rowmap = {}
    ns = sorted(n for n, d in rows.items() if any(k in d for k in cols))
    for n in ns:
        d = rows[n]
        rowmap[n] = r
        ws.cell(r, 1, f"{g['prefix']}{n}")
        names = []
        for key, tc_col in zip(cols, (2, 4, 6)):
            if key in d:
                sec, title = d[key]
                ws.cell(r, tc_col, sec_to_tc(sec))
                ws.cell(r, tc_col + 1, round(sec, 4))
                names.append(title)
            else:
                names.append("—")
        ws.cell(r, 8, f'=IF(AND(E{r}<>"",C{r}<>""),E{r}-C{r},"")')
        ws.cell(r, 9, f'=IF(AND(G{r}<>"",C{r}<>""),G{r}-C{r},"")')
        ws.cell(r, 10, " / ".join(names))
        for c in range(1, 11):
            cell = ws.cell(r, c); cell.border = BORDER
            if c != 10:
                cell.alignment = CENTER
            if c in (3, 5, 7, 8, 9):
                cell.number_format = "0.000"
        r += 1
    LAST = max(r - 1, FIRST)

    sr = LAST + 2
    ws.cell(sr, 1, "统计").font = BOLD
    ws.cell(sr, 2, "有效数量").font = BOLD
    for c in (3, 5, 7, 8, 9):
        L = get_column_letter(c)
        ws.cell(sr, c, f"=COUNT({L}{FIRST}:{L}{LAST})")
    for lab, fn in (("平均(秒)", "AVERAGE"), ("最小(秒)", "MIN"), ("最大(秒)", "MAX")):
        sr += 1
        ws.cell(sr, 2, lab).font = BOLD
        ws.cell(sr, 8, f'=IFERROR({fn}(H{FIRST}:H{LAST}),"")')
        ws.cell(sr, 9, f'=IFERROR({fn}(I{FIRST}:I{LAST}),"")')
    for rr in range(LAST + 2, sr + 1):
        for c in (3, 5, 7, 8, 9):
            ws.cell(rr, c).number_format = "0.000"; ws.cell(rr, c).alignment = CENTER

    for c, w in {1: 7, 2: 13, 3: 10, 4: 13, 5: 10, 6: 13, 7: 10, 8: 9, 9: 9, 10: 34}.items():
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "B3"
    return FIRST, LAST, rowmap


def build_excel(rows, anomalies, ignored, src, out_path):
    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "汇总"
    maps = {}
    for gk in ("start", "end"):
        g = GROUPS[gk]
        ws = wb.create_sheet(g["sheet"])
        maps[gk] = build_group_sheet(ws, g, rows)

    ws = ws_sum
    ws.append(["序号", "句首时延 环境A(b-a)", "句首时延 环境B(c-a)", "句尾时延 环境A(b'-a')", "句尾时延 环境B(c'-a')"])
    for c in range(1, 6):
        cell = ws.cell(1, c); cell.font = BOLD; cell.alignment = CENTER; cell.border = BORDER
        cell.fill = fill("d")
    ns = sorted(rows)
    FIRST = 2
    r = FIRST
    for n in ns:
        ws.cell(r, 1, f"第{n}句")
        for c, (gk, col) in enumerate((("start", "H"), ("start", "I"), ("end", "H"), ("end", "I")), start=2):
            rm = maps[gk][2]
            if n in rm:
                ws.cell(r, c, f"=IF('{GROUPS[gk]['sheet']}'!{col}{rm[n]}=\"\",\"\",'{GROUPS[gk]['sheet']}'!{col}{rm[n]})")
        for c in range(1, 6):
            cell = ws.cell(r, c); cell.border = BORDER; cell.alignment = CENTER
            if c > 1:
                cell.number_format = "0.000"
        r += 1
    LAST = max(r - 1, FIRST)
    ws.cell(r, 1, "平均时延").font = BOLD
    for c in range(2, 6):
        L = get_column_letter(c)
        ws.cell(r, c, f'=IFERROR(AVERAGE({L}{FIRST}:{L}{LAST}),"")')
        ws.cell(r, c).number_format = "0.000"; ws.cell(r, c).alignment = CENTER; ws.cell(r, c).border = BORDER
    ws.cell(r, 1).border = BORDER
    ws.cell(r + 2, 1, f"数据来源：{src}").font = Font(italic=True, color="666666")
    ws.cell(r + 3, 1, f"生成时间：{time.strftime('%Y-%m-%d %H:%M:%S')}；时间码帧率 {FPS}fps；"
                      f"秒 = 剪映标记 start(微秒)/1e6；未识别颜色已忽略 {len(ignored)} 个。").font = Font(italic=True, color="666666")
    for c, w in {1: 9, 2: 20, 3: 20, 4: 20, 5: 20}.items():
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "B2"

    ws2 = wb.create_sheet("检查提示")
    ws2.append(["序号", "提示"])
    ws2["A1"].font = BOLD; ws2["B1"].font = BOLD
    if anomalies:
        for i, a in enumerate(anomalies, 1):
            ws2.append([i, a])
    else:
        ws2.append([1, "无异常：每句标记齐全，无负差值"])
    ws2.column_dimensions["B"].width = 70
    if ignored:
        ws2.append([])
        ws2.append(["已忽略", "颜色 / 时间码 / 名称"])
        for cname, sec, title in sorted(ignored, key=lambda x: x[1]):
            ws2.append(["", f"{cname}  {sec_to_tc(sec)}  {title}"])

    ws3 = wb.create_sheet("原始记录")
    ws3.append(["列", "颜色", "句号", "时间码", "秒", "标记名称"])
    for c in range(1, 7):
        ws3.cell(1, c).font = BOLD
    order = {k: i for i, k in enumerate(KEYS)}
    recs = []
    for n, d in rows.items():
        for key, (sec, title) in d.items():
            recs.append((key, KEYS[key][0], f"{GROUPS[KEYS[key][2]]['prefix']}{n}", sec_to_tc(sec), round(sec, 4), title))
    recs.sort(key=lambda x: (order[x[0]], x[4]))
    for rec in recs:
        ws3.append(list(rec))
    for rr in range(2, len(recs) + 2):
        ws3.cell(rr, 5).number_format = "0.000"
    for c, w in {1: 5, 2: 8, 3: 7, 4: 13, 5: 10, 6: 24}.items():
        ws3.column_dimensions[get_column_letter(c)].width = w

    try:
        wb.save(out_path)
    except PermissionError:
        out_path = out_path.with_name(out_path.stem + time.strftime("_%H%M%S") + ".xlsx")
        wb.save(out_path)
    return out_path


def print_group(g, rows):
    cols = g["cols"]
    ns = sorted(n for n, d in rows.items() if any(k in d for k in cols))
    if not ns:
        print(f"\n[{g['sheet']}] 暂无标记")
        return
    print(f"\n[{g['sheet']}]")
    print(f"{'句':<5}{cols[0]:>12}{cols[1]:>12}{cols[2]:>12}{g['diff'][0]:>8}{g['diff'][1]:>8}")
    for n in ns:
        d = rows[n]
        f = lambda k: sec_to_tc(d[k][0]) if k in d else "—"
        base = cols[0]
        d1 = f"{d[cols[1]][0]-d[base][0]:8.3f}" if base in d and cols[1] in d else f"{'—':>8}"
        d2 = f"{d[cols[2]][0]-d[base][0]:8.3f}" if base in d and cols[2] in d else f"{'—':>8}"
        print(f"{g['prefix']}{n:<4}{f(cols[0]):>12}{f(cols[1]):>12}{f(cols[2]):>12}{d1}{d2}")
